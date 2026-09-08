#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Génère la présentation exécutive POSTE AI (à exposer à la Direction Générale)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

NAVY   = RGBColor(0x0B, 0x1F, 0x3A)
NAVY2  = RGBColor(0x12, 0x2A, 0x4C)
GOLD   = RGBColor(0xC9, 0xA2, 0x27)
INK    = RGBColor(0x1B, 0x24, 0x30)
MUTED  = RGBColor(0x5B, 0x66, 0x75)
LIGHT  = RGBColor(0xF5, 0xF6, 0xF8)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
POS    = RGBColor(0x1E, 0x7F, 0x53)
NEG    = RGBColor(0xC0, 0x39, 0x2B)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height

def slide():
    return prs.slides.add_slide(BLANK)

def rect(s, x, y, w, h, fill, line=None):
    sp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp

def tb(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, sp_after=6, line_sp=1.08):
    box = s.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    if isinstance(runs[0], tuple): runs = [runs]
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(sp_after); p.line_spacing = line_sp
        for (txt, size, bold, color, *rest) in para:
            r = p.add_run(); r.text = txt
            r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
            r.font.name = "Calibri"
            if rest and rest[0] == "italic": r.font.italic = True
    return box

def footer(s, n):
    tb(s, Inches(0.55), SH - Inches(0.5), Inches(9), Inches(0.35),
       [[("POSTE AI  ·  Prototype de démonstration — Données synthétiques — Version conceptuelle", 9, False, MUTED)]])
    tb(s, SW - Inches(1.4), SH - Inches(0.5), Inches(0.9), Inches(0.35),
       [[(str(n), 9, False, MUTED)]], align=PP_ALIGN.RIGHT)

def header(s, kicker, title):
    rect(s, 0, 0, SW, Inches(1.25), NAVY)
    rect(s, 0, Inches(1.25), SW, Pt(3), GOLD)
    tb(s, Inches(0.55), Inches(0.16), Inches(11), Inches(0.35),
       [[(kicker.upper(), 11, True, GOLD)]])
    tb(s, Inches(0.55), Inches(0.5), Inches(12.2), Inches(0.7),
       [[(title, 26, True, WHITE)]])

def content_slide(kicker, title, bullets, note=None, two_col=None):
    s = slide()
    rect(s, 0, 0, SW, SH, LIGHT)
    header(s, kicker, title)
    if two_col:
        left, right = two_col
        _col(s, Inches(0.55), Inches(1.7), Inches(6.0), left)
        _col(s, Inches(6.95), Inches(1.7), Inches(5.8), right)
    else:
        _bullets(s, Inches(0.7), Inches(1.8), Inches(12), bullets)
    if note:
        s.notes_slide.notes_text_frame.text = note
    footer(s, len(prs.slides._sldIdLst))
    return s

def _bullets(s, x, y, w, items):
    paras = []
    for it in items:
        if isinstance(it, tuple):
            lvl, txt = it
        else:
            lvl, txt = 0, it
        if lvl == 0:
            paras.append([("▸  ", 16, True, GOLD), (txt, 16, False, INK)])
        elif lvl == 1:
            paras.append([("      –  ", 13, False, MUTED), (txt, 13, False, MUTED)])
        else:
            paras.append([(txt, 15, True, NAVY)])
    tb(s, x, y, w, Inches(5), paras, sp_after=10, line_sp=1.12)

def _col(s, x, y, w, block):
    title, items = block
    rect(s, x, y, w, Inches(0.5), NAVY)
    tb(s, x + Inches(0.2), y + Inches(0.06), w - Inches(0.3), Inches(0.4),
       [[(title, 13, True, WHITE)]])
    _bullets(s, x + Inches(0.05), y + Inches(0.72), w, items)

# ----------------------------------------------------------------------------
# 1 — Couverture
s = slide()
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, 0, Inches(0.35), SH, GOLD)
# logo mark
lg = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.95), Inches(0.9), Inches(0.9), Inches(0.9))
lg.fill.solid(); lg.fill.fore_color.rgb = GOLD; lg.line.fill.background(); lg.shadow.inherit = False
tf = lg.text_frame; tf.word_wrap = False
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "P"; r.font.size = Pt(34); r.font.bold = True; r.font.color.rgb = NAVY; r.font.name = "Georgia"
tb(s, Inches(2.05), Inches(1.02), Inches(9), Inches(0.7),
   [[("POSTE AI", 30, True, WHITE)]])
tb(s, Inches(2.07), Inches(1.68), Inches(9), Inches(0.4),
   [[("PLATEFORME D'INTELLIGENCE DÉCISIONNELLE", 12, True, GOLD)]])
tb(s, Inches(0.95), Inches(3.15), Inches(11.4), Inches(1.7),
   [[("L'intelligence des données", 44, True, WHITE)],
    [("au service de la décision", 44, True, WHITE)]], line_sp=1.05)
tb(s, Inches(0.95), Inches(5.0), Inches(10.8), Inches(1.0),
   [[("Transformer les données opérationnelles, financières, commerciales et territoriales de "
      "La Poste du Mali en informations exploitables, prévisions et recommandations.", 14, False, RGBColor(0xC9,0xD3,0xE0))]], line_sp=1.2)
tb(s, Inches(0.95), Inches(6.35), Inches(11.4), Inches(0.5),
   [[("Prototype de démonstration — Données synthétiques — À présenter à la Direction Générale", 11, False, GOLD)]])
s.notes_slide.notes_text_frame.text = (
    "Ouverture. POSTE AI n'est pas un chatbot : c'est un système d'intelligence décisionnelle avec "
    "un assistant IA intégré. Objectif de la séance : montrer ce que deviendrait le pilotage de "
    "La Poste une fois connecté aux données réelles.")

# 2 — Le constat
content_slide("Contexte", "Pourquoi maintenant", [
    (2, "Un patrimoine de données riche, mais dispersé"),
    "Les données d'activité, de finance, de logistique et de réseau existent, mais restent "
    "cloisonnées et exploitées manuellement.",
    "Le pilotage repose sur des rapports périodiques, tardifs et hétérogènes.",
    "Les signaux faibles (sous-performance d'agence, dérive de coûts, retards logistiques) "
    "sont détectés trop tard.",
    (2, "Ce que la Direction attend d'un outil de pilotage"),
    "Voir en temps quasi réel, comprendre les causes, anticiper, et disposer de recommandations "
    "actionnables — au même endroit.",
], note="Insister : le problème n'est pas le manque de données, c'est le manque d'un système qui les "
        "transforme en décision.")

# 3 — Ce qu'est POSTE AI
s = content_slide("Le produit", "Un système d'intelligence décisionnelle", [
    (2, "POSTE AI n'est pas un chatbot"),
    "C'est une plateforme de pilotage pour la Direction Générale et les responsables habilités, "
    "avec un assistant IA intégré.",
    "Elle couvre 10 modules : vue générale, performance, agences, courrier & colis, finance, "
    "prévisions, risques & alertes, intelligence IA, scénarios, rapports.",
])
# chaîne de valeur
chain = ["DONNÉES", "ANALYSE", "IA", "PRÉVISION", "RISQUES", "RECOMMANDATION", "DÉCISION"]
x = Inches(0.6); y = Inches(5.05); bw = Inches(1.62); gap = Inches(0.14)
for i, c in enumerate(chain):
    col = GOLD if c == "DÉCISION" else NAVY
    b = rect(s, x + i*(bw+gap), y, bw, Inches(0.75), col)
    tf = b.text_frame; tf.word_wrap = True; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = c; r.font.size = Pt(10); r.font.bold = True
    r.font.color.rgb = NAVY if c == "DÉCISION" else WHITE; r.font.name = "Calibri"
    if i < len(chain) - 1:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x + i*(bw+gap) + bw - Inches(0.02), y + Inches(0.27),
                                Inches(0.18), Inches(0.2))
        ar.fill.solid(); ar.fill.fore_color.rgb = GOLD; ar.line.fill.background(); ar.shadow.inherit = False
tb(s, Inches(0.6), Inches(5.95), Inches(12), Inches(0.4),
   [[("De la donnée brute à la décision informée — une seule chaîne, un seul outil.", 12, False, MUTED, "italic")]])

# 4 — Pour qui
content_slide("Utilisateurs", "Conçu pour les décideurs", bullets=None, two_col=(
    ("Directions", [
        "Direction Générale",
        "Direction des opérations",
        "Direction financière",
        "Direction commerciale",
        "Direction informatique",
    ]),
    ("Réseau & analyse", [
        "Responsables régionaux",
        "Responsables d'agences",
        "Analystes / contrôle de gestion",
        (1, "Accès par rôle (RBAC) : ADMIN, DIRECTION_GÉNÉRALE, DIRECTION_FINANCIÈRE, "
            "DIRECTION_OPÉRATIONS, RESPONSABLE_RÉGIONAL, ANALYSTE"),
    ]),
), note="Chaque rôle voit un périmètre adapté. L'interface est déjà pensée pour ce cloisonnement.")

# 5 — Philosophie produit
s = slide()
rect(s, 0, 0, SW, SH, LIGHT)
header(s, "Philosophie", "SEE · UNDERSTAND · PREDICT · RECOMMEND · DECIDE")
steps = [
    ("SEE", "Voir ce qui se passe", "KPI consolidés, performance par région et par agence."),
    ("UNDERSTAND", "Comprendre pourquoi", "AI Insights : causes, contributions, tendances."),
    ("PREDICT", "Anticiper", "Prévisions 30 / 90 jours avec intervalle de confiance."),
    ("RECOMMEND", "Recommander", "Actions prioritaires générées et traçables."),
    ("DECIDE", "Décider", "Simulateur de scénarios avant l'arbitrage."),
]
y = Inches(1.75)
for i, (k, t, d) in enumerate(steps):
    yy = y + Inches(i * 1.02)
    n = rect(s, Inches(0.7), yy, Inches(1.9), Inches(0.8), NAVY)
    tf = n.text_frame; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = k; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = GOLD; r.font.name="Calibri"
    tb(s, Inches(2.9), yy + Inches(0.02), Inches(9.8), Inches(0.8),
       [[(t + " — ", 15, True, NAVY), (d, 13, False, MUTED)]], anchor=MSO_ANCHOR.MIDDLE)
footer(s, 5)
s.notes_slide.notes_text_frame.text = ("Ces 5 verbes structurent toute l'application et le déroulé de la démo.")

# 6..14 — modules (capture d'écran optionnelle à insérer plus tard)
modules = [
 ("Module 01", "Vue générale — Tableau de bord exécutif", [
    "KPI consolidés : chiffre d'affaires, colis, transactions, courrier, agences actives, taux de service.",
    "Graphique d'évolution multi-métriques + performance par région avec badges (Excellent → Critique).",
    "AI Insights automatiques et panneau d'alertes prioritaires.",
    (1, "En < 10 s : où en est La Poste, ce qui va bien, ce qui pose problème, ce qui risque d'arriver."),
 ], "Écran d'ouverture de la démo : « Voici la situation globale de La Poste »."),
 ("Modules 02–03", "Performance & Réseau d'agences", [
    "Analyse par CA, transactions, colis, courrier, dépenses, agents, fréquentation — filtres région / agence / service.",
    "Table de 30 agences : score, croissance, coût, rentabilité, risque, statut.",
    "Fiche agence : score /100, historique, diagnostic IA et recommandations ciblées.",
    (1, "« Nous pouvons descendre jusqu'au niveau opérationnel. »"),
 ], "Cliquer une agence en difficulté et lire le diagnostic IA."),
 ("Module 04", "Courrier & Colis — Intelligence logistique", [
    "Colis entrants / sortants, livraisons, retards, taux de livraison, délai moyen.",
    "Volumes et délais sur 12 mois, répartition par région, Top 10 des corridors.",
    "Alertes logistiques (ex. corridor Nord : délais anormaux).",
 ], None),
 ("Module 05", "Finance — Intelligence financière", [
    "Revenus, dépenses, marge, coûts opérationnels, budget, écart budgétaire.",
    "Revenus vs dépenses, budget vs réalisé, rentabilité par agence.",
    "AI Financial Insights : « les coûts opérationnels progressent plus vite que les revenus dans certaines agences ».",
 ], None),
 ("Module 06", "Prévisions — Forecasting & Predictive Intelligence", [
    "Projection du CA, des colis, des transactions et du courrier à 30 et 90 jours.",
    "Intervalle de confiance à 90 % et indicateur de fiabilité du modèle.",
    (1, "Libellé explicite : MODÈLE SIMULÉ — DONNÉES DE DÉMONSTRATION."),
 ], "« Nous pouvons anticiper l'évolution de l'activité. »"),
 ("Module 07", "Risques & Anomalies", [
    "Matrice probabilité × impact par catégorie (opérationnel, financier, commercial, logistique, réseau).",
    "Détection d'anomalies : valeur attendue vs observée, écart, sévérité.",
    "Alertes prioritaires : date, niveau, objet, impact, statut.",
 ], "« Le système détecte automatiquement un comportement inhabituel. »"),
 ("Module 08", "POSTE AI Copilot — Intelligence IA", [
    "Questions en langage naturel : « Quelles agences présentent les plus grands risques ? »",
    "Réponses analytiques : synthèse chiffrée, graphiques intégrés, causes probables, recommandations.",
    (1, "Aide à la décision — pas de décision automatique. Traçabilité et validation humaine."),
 ], "Demander : « Quelles sont les principales priorités pour la Direction ? »"),
 ("Module 09", "Decision Simulator — Scénarios", [
    "Leviers : investissement, personnel, horaires, capacité logistique, digitalisation.",
    "Impact estimé sur CA, transactions, colis, coûts et rentabilité, en temps réel.",
    (1, "Simulation indicative basée sur des données synthétiques."),
 ], "« Nous pouvons tester différents scénarios avant de prendre une décision. »"),
 ("Module 10", "Rapports exécutifs", [
    "Rapports quotidien, hebdomadaire, mensuel, trimestriel.",
    "Aperçu, génération et export PDF (simulés dans le prototype).",
    "Centre de notifications : alertes, anomalies, prévisions, rapports disponibles.",
 ], None),
]
for i, (k, t, b, note) in enumerate(modules):
    content_slide(k, t, b, note=note)

# 15 — Architecture & sécurité
content_slide("Sous le capot", "Architecture cible & sécurité", bullets=None, two_col=(
    ("Chaîne de traitement (Phase 2+)", [
        "Sources → Data Integration Layer",
        "Data Lake / Data Warehouse (PostgreSQL)",
        "Analytics Engine (KPI, scores, classements)",
        "ML Engine (prévision, anomalies, risques)",
        "AI / LLM Layer (copilote, RAG, explications)",
        "Decision Engine → Dashboard · Copilot · Reports",
    ]),
    ("Sécurité & IA responsable", [
        "RBAC, moindre privilège, journaux d'audit",
        "Chiffrement, authentification des API, gouvernance",
        "Traçabilité des recommandations, niveau de confiance",
        "Explication des prédictions, validation humaine",
        (1, "Aucune décision critique automatisée sans contrôle humain — DECISION SUPPORT, pas AUTOMATED DECISION MAKER."),
    ]),
), note="Rassurer la DSI : rien n'est figé, l'architecture est standard et sécurisée par conception.")

# 16 — Roadmap
s = slide()
rect(s, 0, 0, SW, SH, LIGHT)
header(s, "Trajectoire", "Vision de déploiement")
phases = [
    ("Phase 1", "Prototype", "Maquette sur données synthétiques — validation de la valeur (en cours)."),
    ("Phase 2", "Pilote données réelles", "Un périmètre pilote + Data Warehouse + premiers modèles ML."),
    ("Phase 3", "Connexion aux SI", "Intégration des applications métier, finance, logistique, réseau."),
    ("Phase 4", "Déploiement institutionnel", "Généralisation, formation, exploitation en continu."),
    ("Phase 5", "Prédictif avancé", "Recommandations proactives, LLM + RAG sur le patrimoine de données."),
]
y = Inches(1.8)
for i, (p, t, d) in enumerate(phases):
    yy = y + Inches(i * 1.02)
    col = GOLD if i == 0 else NAVY
    n = rect(s, Inches(0.7), yy, Inches(1.75), Inches(0.82), col)
    tf = n.text_frame; par = tf.paragraphs[0]; par.alignment = PP_ALIGN.CENTER
    r = par.add_run(); r.text = p; r.font.size = Pt(12); r.font.bold = True
    r.font.color.rgb = NAVY if i == 0 else WHITE; r.font.name = "Calibri"
    tb(s, Inches(2.75), yy + Inches(0.03), Inches(10), Inches(0.8),
       [[(t + " — ", 15, True, NAVY), (d, 12.5, False, MUTED)]], anchor=MSO_ANCHOR.MIDDLE)
footer(s, len(prs.slides._sldIdLst))

# 17 — Prochaines étapes / demande
content_slide("Décision attendue", "Prochaines étapes", [
    (2, "Ce que nous proposons"),
    "Valider le principe d'un système d'intelligence décisionnelle pour La Poste du Mali.",
    "Désigner un périmètre pilote (une direction ou une région) et un référent métier.",
    "Donner accès à un premier jeu de données réelles pour la Phase 2.",
    (2, "Ce que la Phase 2 démontrera"),
    "Les mêmes écrans, alimentés par les vrais chiffres, avec prévisions recalées quotidiennement.",
    "Un cadre de gouvernance et de sécurité opérationnel.",
], note="Formuler une demande claire et limitée : un périmètre pilote + un accès données.")

# 18 — Clôture
s = slide()
rect(s, 0, 0, SW, SH, NAVY)
rect(s, 0, 0, Inches(0.35), SH, GOLD)
tb(s, Inches(0.95), Inches(2.6), Inches(11.4), Inches(1.4),
   [[("« Voici à quoi pourrait ressembler le futur", 30, True, WHITE)],
    [("système d'intelligence décisionnelle de La Poste du Mali. »", 30, True, WHITE)]], line_sp=1.1)
tb(s, Inches(0.98), Inches(4.5), Inches(11), Inches(0.5),
   [[("POSTE AI — SEE · UNDERSTAND · PREDICT · RECOMMEND · DECIDE", 13, True, GOLD)]])
tb(s, Inches(0.98), Inches(6.4), Inches(11.4), Inches(0.5),
   [[("Prototype de démonstration — Données synthétiques — Aucune donnée réelle de La Poste du Mali.", 10, False, RGBColor(0xC9,0xD3,0xE0))]])

out = "/home/anonymous/Documents/AI_WORK/POSTE_AI/POSTE_AI_presentation_DG.pptx"
prs.save(out)
print("OK ->", out, "|", len(prs.slides._sldIdLst), "slides")
